import os
import tempfile
from openpyxl import load_workbook
from docx import Document

async def convert_excel_to_word(file) -> str:
    with tempfile.TemporaryDirectory() as tmpdirname:
        excel_path = os.path.join(tmpdirname, file.filename)
        
        # Save uploaded Excel file
        with open(excel_path, "wb") as f:
            f.write(await file.read())

        # Load Excel workbook
        wb = load_workbook(excel_path)
        ws = wb.active

        # Create a new Word document
        doc = Document()
        doc.add_heading("Excel to Word Conversion", level=1)

        # Write Excel data to Word
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join([str(cell) if cell is not None else "" for cell in row])
            doc.add_paragraph(row_text)

        # Save Word file in static folder
        output_path = "static/output.docx"
        os.makedirs("static", exist_ok=True)
        doc.save(output_path)

        return output_path
