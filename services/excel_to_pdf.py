import os
import tempfile
from openpyxl import load_workbook
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from fastapi import UploadFile

async def convert_excel_to_pdf(file: UploadFile) -> str:
    with tempfile.TemporaryDirectory() as tmpdirname:
        # Save uploaded Excel file
        excel_path = os.path.join(tmpdirname, file.filename)
        with open(excel_path, "wb") as f:
            f.write(await file.read())

        # Load workbook and sheet
        wb = load_workbook(excel_path)
        ws = wb.active

        # Extract data
        data = []
        for row in ws.iter_rows(values_only=True):
            data.append(list(row))

        # Output PDF path
        pdf_path = os.path.join(tmpdirname, "output.pdf")
        doc = SimpleDocTemplate(pdf_path, pagesize=letter)

        # Create table
        table = Table(data)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.gray),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ]))

        elements = [table]
        doc.build(elements)

        # Save PDF to static folder
        output_path = "static/excel_output.pdf"
        os.makedirs("static", exist_ok=True)
        os.replace(pdf_path, output_path)

        return output_path
